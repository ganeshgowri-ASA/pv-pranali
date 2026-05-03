"""Vidyalaya MCP server — wraps vidyalaya-office document generation.

Tools: create_docx, create_xlsx, create_pptx, export_pdf, parse_docx

Usage:
    python mcp_servers/vidyalaya/server.py
"""
from __future__ import annotations

import base64
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from mcp_servers._lib import MCPClient

try:
    from docx import Document
    from docx.shared import Pt
except ImportError:
    Document = None  # type: ignore[assignment]
    Pt = None  # type: ignore[assignment]

try:
    import openpyxl
except ImportError:
    openpyxl = None  # type: ignore[assignment]

try:
    from pptx import Presentation
    from pptx.util import Inches
except ImportError:
    Presentation = None  # type: ignore[assignment]
    Inches = None  # type: ignore[assignment]

mcp = MCPClient("vidyalaya")

_TMP = Path(tempfile.gettempdir()) / "vidyalaya"
_TMP.mkdir(exist_ok=True)


def _to_base64(path: Path) -> str:
    return base64.b64encode(path.read_bytes()).decode()


@mcp.tool()
def create_docx(
    title: str,
    paragraphs: list[str],
    output_path: str = "",
) -> dict[str, Any]:
    """Create a .docx Word document with a title and body paragraphs.

    Args:
        title: Document title, written as Heading 1.
        paragraphs: List of paragraph strings to append.
        output_path: Optional absolute path for the output .docx file.

    Returns:
        {"path": str, "base64": str} or {"error": str}
    """
    if Document is None:
        return {"error": "python-docx not installed — run: pip install python-docx"}
    doc = Document()
    doc.add_heading(title, level=1)
    for para in paragraphs:
        doc.add_paragraph(para)
    dest = (
        Path(output_path)
        if output_path
        else _TMP / f"{title[:40].replace(' ', '_')}.docx"
    )
    doc.save(dest)
    return {"path": str(dest), "base64": _to_base64(dest)}


@mcp.tool()
def create_xlsx(
    sheet_name: str,
    rows: list[list],
    headers: list[str] | None = None,
    output_path: str = "",
) -> dict[str, Any]:
    """Create a .xlsx spreadsheet with a named sheet, optional headers, and rows.

    Args:
        sheet_name: Worksheet tab name.
        rows: List of rows; each row is a list of cell values.
        headers: Optional column headers prepended before data rows.
        output_path: Optional absolute path for the output .xlsx file.

    Returns:
        {"path": str, "base64": str} or {"error": str}
    """
    if openpyxl is None:
        return {"error": "openpyxl not installed — run: pip install openpyxl"}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    if headers:
        ws.append(headers)
    for row in rows:
        ws.append(row)
    dest = (
        Path(output_path)
        if output_path
        else _TMP / f"{sheet_name[:40].replace(' ', '_')}.xlsx"
    )
    wb.save(dest)
    return {"path": str(dest), "base64": _to_base64(dest)}


@mcp.tool()
def create_pptx(
    title: str,
    slides: list[dict],
    output_path: str = "",
) -> dict[str, Any]:
    """Create a .pptx presentation with a cover slide and content slides.

    Args:
        title: Presentation title shown on the cover slide.
        slides: List of slide dicts, each with 'title' (str) and optional 'body' (str).
        output_path: Optional absolute path for the output .pptx file.

    Returns:
        {"path": str, "base64": str} or {"error": str}
    """
    if Presentation is None:
        return {"error": "python-pptx not installed — run: pip install python-pptx"}
    prs = Presentation()
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = title
    if len(cover.placeholders) > 1:
        cover.placeholders[1].text = ""
    content_layout = prs.slide_layouts[1]
    for s in slides:
        sl = prs.slides.add_slide(content_layout)
        sl.shapes.title.text = s.get("title", "")
        if len(sl.placeholders) > 1:
            sl.placeholders[1].text = s.get("body", "")
    dest = (
        Path(output_path)
        if output_path
        else _TMP / f"{title[:40].replace(' ', '_')}.pptx"
    )
    prs.save(dest)
    return {"path": str(dest), "base64": _to_base64(dest)}


@mcp.tool()
def export_pdf(
    input_path: str,
    output_dir: str = "",
) -> dict[str, Any]:
    """Convert a .docx or .pptx file to PDF using LibreOffice headless.

    Args:
        input_path: Absolute path to the source .docx or .pptx file.
        output_dir: Optional directory for the PDF (defaults to same dir as input).

    Returns:
        {"path": str} or {"error": str}
    """
    src = Path(input_path)
    if not src.exists():
        return {"error": f"File not found: {input_path}"}
    out_dir = Path(output_dir) if output_dir else src.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    cmd = [
        "libreoffice",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(out_dir),
        str(src),
    ]
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    except FileNotFoundError:
        return {"error": "LibreOffice not found — install with: apt install libreoffice"}
    if result.returncode != 0:
        return {"error": result.stderr or "LibreOffice conversion failed"}
    pdf_path = out_dir / (src.stem + ".pdf")
    return {"path": str(pdf_path)}


@mcp.tool()
def parse_docx(file_path: str) -> dict[str, Any]:
    """Parse a .docx file and extract its title and paragraph text.

    Args:
        file_path: Absolute path to the .docx file.

    Returns:
        {"title": str, "paragraphs": list[str], "path": str} or {"error": str}
    """
    if Document is None:
        return {"error": "python-docx not installed — run: pip install python-docx"}
    src = Path(file_path)
    if not src.exists():
        return {"error": f"File not found: {file_path}"}
    doc = Document(str(src))
    title = ""
    paragraphs: list[str] = []
    for para in doc.paragraphs:
        if not title and para.style.name.startswith("Heading"):
            title = para.text
        elif para.text.strip():
            paragraphs.append(para.text)
    return {"title": title, "paragraphs": paragraphs, "path": str(src)}


if __name__ == "__main__":
    mcp.run()
